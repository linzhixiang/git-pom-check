import io
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import uuid
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from urllib.parse import urlparse, quote

from fastapi import FastAPI, UploadFile, File, HTTPException, Header
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from pydantic import BaseModel

from .analyzer import (
    collect_maven, collect_gradle, deduplicate_nodes,
    clone_gitlab, detect_build_system, resolve_maven, resolve_git,
    resolve_m2_repos, compute_jar_sha256,
)


# ========== 内置 GitLab Token（改这里即可） ==========
DEFAULT_GITLAB_TOKEN = "glpat-tQFnAx33P2rjYKyCq7xs"   # ← 替换成真实 token


TASK_EXECUTOR = ThreadPoolExecutor(max_workers=2)
PROJECT_POOL_SIZE = 4
MAX_BATCH_PROJECTS = 30
TASKS: Dict[str, Dict[str, Any]] = {}
TASKS_LOCK = threading.Lock()
TASK_TTL_SECONDS = 3600
MAX_UPLOAD_BYTES = 50 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


# ========== 项目名 / 文件名 / sheet 名 处理 ==========
def extract_project_name(repo_url: str) -> str:
    try:
        parsed = urlparse(repo_url)
        path = parsed.path.strip("/")
        if not path:
            return "unknown_project"
        last = path.split("/")[-1]
        if last.endswith(".git"):
            last = last[:-4]
        return last or "unknown_project"
    except Exception:
        return "unknown_project"


def sanitize_filename(name: str, fallback: str = "project") -> str:
    if not name:
        return fallback
    cleaned = re.sub(r'[\\/:*?"<>|\r\n\t]', "_", name).strip().strip(". ")
    if not cleaned:
        return fallback
    return cleaned[:100]


def sanitize_sheet_name(name: str, used: set) -> str:
    cleaned = re.sub(r'[:\\/?*\[\]]', "_", name or "Sheet")
    cleaned = cleaned.strip().strip("'")
    if not cleaned:
        cleaned = "Sheet"
    if len(cleaned) > 31:
        cleaned = cleaned[:31]
    base = cleaned
    i = 1
    while cleaned.lower() in used:
        suffix = f"_{i}"
        cleaned = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(cleaned.lower())
    return cleaned


def build_download_filename(projects: List[Dict[str, Any]]) -> str:
    if len(projects) == 1:
        pname = sanitize_filename(projects[0].get("project_name") or "project", "project")
        return f"{pname}_开源软件（源码）资产台账"
    return f"开源软件资产台账_{len(projects)}个项目"


# ========== GitLab Token 解析 ==========
def resolve_gitlab_token(body_token: Optional[str], header_token: Optional[str]) -> Tuple[Optional[str], str]:
    if body_token and body_token.strip():
        return body_token.strip(), "body"
    if header_token and header_token.strip():
        return header_token.strip(), "header"
    env_token = os.environ.get("GITLAB_TOKEN", "").strip()
    if env_token:
        return env_token, "env"
    if DEFAULT_GITLAB_TOKEN and DEFAULT_GITLAB_TOKEN.strip():
        return DEFAULT_GITLAB_TOKEN.strip(), "builtin"
    return None, "none"


# ========== 任务管理 ==========
def cleanup_expired() -> None:
    now = time.time()
    with TASKS_LOCK:
        expired = []
        for tid, t in TASKS.items():
            finished = t.get("finished_at")
            if not finished:
                continue
            try:
                if now - datetime.fromisoformat(finished).timestamp() > TASK_TTL_SECONDS:
                    expired.append(tid)
            except ValueError:
                continue
        for tid in expired:
            del TASKS[tid]


# ========== 安全解压 ==========
def safe_extract_zip(zip_path: Path, target_dir: Path) -> None:
    target_root = target_dir.resolve()
    total = 0
    with zipfile.ZipFile(zip_path) as zf:
        for info in zf.infolist():
            name = info.filename
            if name.endswith("/"):
                continue
            out_path = (target_root / name).resolve()
            if not (str(out_path) == str(target_root)
                    or str(out_path).startswith(str(target_root) + os.sep)):
                raise ValueError(f"非法路径: {name}")
            total += info.file_size
            if total > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("解压后总大小超限")
        zf.extractall(target_dir)


# ========== 单项目解析结果生成 ==========
def analyze_project_dir(project_dir: Path) -> Dict[str, Any]:
    build_system = detect_build_system(project_dir)
    if build_system == "maven":
        nodes, warnings, skipped = collect_maven(project_dir)
    elif build_system == "gradle":
        nodes, warnings, skipped = collect_gradle(project_dir)
    else:
        raise ValueError("未找到 pom.xml 或 build.gradle")

    if not nodes:
        msg = "未解析到任何依赖"
        if warnings:
            msg += "；诊断信息: " + " | ".join(warnings)
        raise ValueError(msg)

    unique = deduplicate_nodes(nodes)
    direct = [n for n in unique if n.depth == 1]
    transitive = [n for n in unique if n.depth > 1]

    deps_with_sha: List[Dict[str, Any]] = []
    for n in unique:
        d = n.to_dict()
        d["sha256"] = compute_jar_sha256(n.group, n.artifact, n.version) or None
        deps_with_sha.append(d)

    is_fallback = any(
        ("降级解析" in w) or ("未生成 dependency-tree.json" in w)
        for w in warnings
    )
    fallback_reason = None
    if is_fallback:
        for w in warnings:
            if "降级解析" in w or "未生成 dependency-tree.json" in w:
                fallback_reason = w
                break

    return {
        "build_system": build_system,
        "summary": {
            "total_unique": len(unique),
            "direct_count": len(direct),
            "transitive_count": len(transitive),
            "skipped_count": len(skipped),
        },
        "dependencies": sorted(
            deps_with_sha,
            key=lambda x: (x["depth"], x["group"], x["artifact"]),
        ),
        "skipped_dependencies": skipped,
        "warnings": warnings,
        "is_fallback": is_fallback,
        "fallback_reason": fallback_reason,
    }


# ========== Excel 构建（1 或 N 个项目） ==========
def _fill_dep_sheet(ws, project_entry: Dict[str, Any], include_transitive: bool) -> None:
    ws["A1"] = "组件的官方规范名"
    ws["B1"] = "实际引入的精确版本号"
    ws["C1"] = "一般内部已自带.sha1/.sha256文件"
    ws["D1"] = "引入层级（直接/间接）"
    ws["A2"] = "组件名称"
    ws["B2"] = "组件版本"
    ws["C2"] = "组件哈希值"
    ws["D2"] = "依赖类型"

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    center = Alignment(horizontal="center", vertical="center")
    for row in (1, 2):
        for col in ("A", "B", "C", "D"):
            cell = ws[f"{col}{row}"]
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center

    if project_entry["status"] != "success" or not project_entry.get("result"):
        err = project_entry.get("error") or "项目解析未成功完成"
        ws["A3"] = f"[项目解析失败] {err}"
        ws.merge_cells("A3:D3")
        ws["A3"].fill = PatternFill("solid", fgColor="FFF2CC")
        ws.column_dimensions["A"].width = 40
        ws.column_dimensions["B"].width = 25
        ws.column_dimensions["C"].width = 80
        ws.column_dimensions["D"].width = 15
        return

    deps = project_entry["result"].get("dependencies", [])
    if not include_transitive:
        deps = [d for d in deps if not d.get("is_transitive")]

    deps_sorted = sorted(
        deps,
        key=lambda x: (
            0 if not x.get("is_transitive") else 1,
            x.get("group") or "",
            x.get("artifact") or "",
        ),
    )

    for i, d in enumerate(deps_sorted, start=3):
        ws[f"A{i}"] = d.get("artifact") or d.get("group") or ""
        ws[f"B{i}"] = d.get("version") or ""
        ws[f"C{i}"] = d.get("sha256") or ""
        ws[f"D{i}"] = "间接依赖" if d.get("is_transitive") else "直接依赖"

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 80
    ws.column_dimensions["D"].width = 15
    ws.freeze_panes = "A3"

    # 如果有被跳过的依赖，追加在依赖表尾部
    skipped = project_entry["result"].get("skipped_dependencies") or []
    if skipped:
        start = len(deps_sorted) + 4
        ws[f"A{start}"] = "【以下依赖无法从远程仓库下载，已被跳过（未纳入依赖树）】"
        ws[f"A{start}"].font = Font(bold=True, color="C00000")
        ws.merge_cells(f"A{start}:D{start}")
        for j, s in enumerate(skipped, start=start + 1):
            ws[f"A{j}"] = s.get("artifact") or s.get("group") or ""
            ws[f"B{j}"] = s.get("version") or ""
            ws[f"C{j}"] = ""
            ws[f"D{j}"] = "跳过（无法下载）"
            ws[f"D{j}"].fill = PatternFill("solid", fgColor="FCE4EC")


def build_excel(projects: List[Dict[str, Any]], include_transitive: bool = True) -> bytes:
    wb = Workbook()
    default_ws = wb.active
    wb.remove(default_ws)

    used_names: set = set()

    if len(projects) > 1:
        ws_sum = wb.create_sheet("汇总")
        used_names.add("汇总")
        headers = ["序号", "项目名", "仓库地址", "状态", "依赖总数", "直接依赖", "传递依赖", "跳过数", "错误信息"]
        header_font = Font(bold=True)
        header_fill = PatternFill("solid", fgColor="DDEBF7")
        center = Alignment(horizontal="center", vertical="center")
        for col, h in enumerate(headers, 1):
            cell = ws_sum.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center

        for i, p in enumerate(projects, 1):
            if p["status"] == "success" and p.get("result"):
                s = p["result"]["summary"]
                total = s["total_unique"]
                direct = s["direct_count"]
                transitive = s["transitive_count"]
                skipped_n = s.get("skipped_count", 0)
            else:
                total = direct = transitive = skipped_n = ""
            ws_sum.cell(row=i + 1, column=1, value=i)
            ws_sum.cell(row=i + 1, column=2, value=p.get("project_name") or "")
            ws_sum.cell(row=i + 1, column=3, value=p.get("repo_url") or "")
            ws_sum.cell(row=i + 1, column=4, value=p.get("status") or "")
            ws_sum.cell(row=i + 1, column=5, value=total)
            ws_sum.cell(row=i + 1, column=6, value=direct)
            ws_sum.cell(row=i + 1, column=7, value=transitive)
            ws_sum.cell(row=i + 1, column=8, value=skipped_n)
            ws_sum.cell(row=i + 1, column=9, value=p.get("error") or "")
            if p["status"] == "failed":
                for col in range(1, 10):
                    ws_sum.cell(row=i + 1, column=col).fill = PatternFill("solid", fgColor="FCE4EC")

        for letter, width in [("A", 6), ("B", 30), ("C", 60), ("D", 12),
                              ("E", 12), ("F", 12), ("G", 12), ("H", 10), ("I", 60)]:
            ws_sum.column_dimensions[letter].width = width
        ws_sum.freeze_panes = "A2"

    for p in projects:
        sheet_name = sanitize_sheet_name(p.get("project_name") or "project", used_names)
        ws = wb.create_sheet(sheet_name)
        _fill_dep_sheet(ws, p, include_transitive)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ========== 上传任务执行 ==========
def run_upload_task(task_id: str, work_dir: Path, extract_dir: Path, project_name: str) -> None:
    with TASKS_LOCK:
        TASKS[task_id]["status"] = "running"
        TASKS[task_id]["started_at"] = now_iso()
        if TASKS[task_id]["projects"]:
            TASKS[task_id]["projects"][0]["status"] = "running"
            TASKS[task_id]["projects"][0]["started_at"] = now_iso()

    try:
        result = analyze_project_dir(extract_dir)
        with TASKS_LOCK:
            p = TASKS[task_id]["projects"][0]
            p["result"] = {
                k: result[k] for k in (
                    "build_system", "summary", "dependencies",
                    "skipped_dependencies", "is_fallback", "fallback_reason",
                )
            }
            p["warnings"] = result["warnings"]
            p["status"] = "success"
            p["finished_at"] = now_iso()
            TASKS[task_id]["summary"] = {"total": 1, "success": 1, "failed": 0}
            TASKS[task_id]["status"] = "success"
            TASKS[task_id]["finished_at"] = now_iso()
    except Exception as e:
        with TASKS_LOCK:
            p = TASKS[task_id]["projects"][0]
            p["status"] = "failed"
            p["error"] = str(e)
            p["finished_at"] = now_iso()
            TASKS[task_id]["summary"] = {"total": 1, "success": 0, "failed": 1}
            TASKS[task_id]["status"] = "failed"
            TASKS[task_id]["finished_at"] = now_iso()
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


# ========== 批量 GitLab 任务执行 ==========
def run_batch_task(task_id: str, work_root: Path, token_map: Dict[int, Optional[str]]) -> None:
    with TASKS_LOCK:
        TASKS[task_id]["status"] = "running"
        TASKS[task_id]["started_at"] = now_iso()
        project_count = len(TASKS[task_id]["projects"])

    def worker(index: int) -> None:
        with TASKS_LOCK:
            p = TASKS[task_id]["projects"][index]
            p["status"] = "running"
            p["started_at"] = now_iso()

        proj_dir = work_root / f"p{index}"
        proj_dir.mkdir(parents=True, exist_ok=True)
        warnings: List[str] = []

        try:
            clone_dir = proj_dir / "repo"
            clone_gitlab(p["repo_url"], token_map.get(index), p.get("branch"), clone_dir)
            if p.get("subpath"):
                target = (clone_dir / p["subpath"]).resolve()
                if not str(target).startswith(str(clone_dir.resolve()) + os.sep):
                    raise ValueError("subpath 越界")
                if not target.exists():
                    raise ValueError(f"subpath 不存在: {p['subpath']}")
                clone_dir = target

            result = analyze_project_dir(clone_dir)
            warnings = result.get("warnings", [])

            with TASKS_LOCK:
                p["result"] = {
                    k: result[k] for k in (
                        "build_system", "summary", "dependencies",
                        "skipped_dependencies", "is_fallback", "fallback_reason",
                    )
                }
                p["warnings"] = warnings
                p["status"] = "success"
                p["finished_at"] = now_iso()
        except Exception as e:
            with TASKS_LOCK:
                p["status"] = "failed"
                p["error"] = str(e)
                p["warnings"] = warnings
                p["finished_at"] = now_iso()
        finally:
            shutil.rmtree(proj_dir, ignore_errors=True)

    try:
        workers = max(1, min(PROJECT_POOL_SIZE, project_count))
        with ThreadPoolExecutor(max_workers=workers) as pool:
            list(pool.map(worker, range(project_count)))
    finally:
        with TASKS_LOCK:
            projects = TASKS[task_id]["projects"]
            success = sum(1 for p in projects if p["status"] == "success")
            failed = sum(1 for p in projects if p["status"] == "failed")
            total = len(projects)
            if success == total:
                overall = "success"
            elif failed == total:
                overall = "failed"
            else:
                overall = "partial"
            TASKS[task_id].update({
                "status": overall,
                "finished_at": now_iso(),
                "summary": {"total": total, "success": success, "failed": failed},
            })
        shutil.rmtree(work_root, ignore_errors=True)


# ========== 创建 GitLab 任务 ==========
def _create_gitlab_task(
    projects_req: List["GitLabRequest"],
    x_gitlab_token: Optional[str],
    batch_mode: bool,
) -> Dict[str, Any]:
    cleanup_expired()

    if not projects_req:
        raise HTTPException(400, "projects 不能为空")
    if len(projects_req) > MAX_BATCH_PROJECTS:
        raise HTTPException(400, f"单次最多 {MAX_BATCH_PROJECTS} 个项目")

    task_id = uuid.uuid4().hex
    work_root = Path(tempfile.mkdtemp(prefix="depgov-gitlab-"))
    token_map: Dict[int, Optional[str]] = {}
    project_entries: List[Dict[str, Any]] = []

    for i, proj in enumerate(projects_req):
        token, source = resolve_gitlab_token(proj.token, x_gitlab_token)
        token_map[i] = token
        pname = sanitize_filename(
            extract_project_name(proj.repo_url), fallback=f"project_{i+1}"
        )
        project_entries.append({
            "index": i,
            "repo_url": proj.repo_url,
            "branch": proj.branch,
            "subpath": proj.subpath,
            "project_name": pname,
            "token_source": source,
            "status": "pending",
            "started_at": None,
            "finished_at": None,
            "result": None,
            "error": None,
            "warnings": [],
        })

    with TASKS_LOCK:
        TASKS[task_id] = {
            "task_id": task_id,
            "source": "gitlab_batch" if batch_mode else "gitlab",
            "batch": batch_mode,
            "status": "pending",
            "created_at": now_iso(),
            "started_at": None,
            "finished_at": None,
            "projects": project_entries,
            "summary": {"total": len(project_entries), "success": 0, "failed": 0},
        }

    TASK_EXECUTOR.submit(run_batch_task, task_id, work_root, token_map)

    return {
        "task_id": task_id,
        "status": "pending",
        "project_count": len(project_entries),
        "task_url": f"/api/v1/tasks/{task_id}",
        "excel_url": f"/api/v1/tasks/{task_id}/excel",
    }


# ========== FastAPI 应用 ==========
app = FastAPI(title="Java Dependency Governance API", version="2.2.0")


class GitLabRequest(BaseModel):
    repo_url: str
    token: Optional[str] = None
    branch: Optional[str] = None
    subpath: Optional[str] = None


class GitLabBatchRequest(BaseModel):
    projects: List[GitLabRequest]


@app.get("/api/v1/health")
def health():
    return {"status": "ok"}


@app.get("/api/v1/diagnostics")
def diagnostics():
    mvn = resolve_maven()
    git = resolve_git()
    _, default_source = resolve_gitlab_token(None, None)
    info: Dict[str, Any] = {
        "platform": sys.platform,
        "executable": sys.executable,
        "PATH": os.environ.get("PATH", ""),
        "MAVEN_BIN_env": os.environ.get("MAVEN_BIN"),
        "MAVEN_WRAPPER_DISABLE": os.environ.get("MAVEN_WRAPPER_DISABLE", "true"),
        "MAVEN_STUB_RETRY": os.environ.get("MAVEN_STUB_RETRY", "true"),
        "MAVEN_STUB_MAX_RETRIES": os.environ.get("MAVEN_STUB_MAX_RETRIES", "3"),
        "GIT_BIN_env": os.environ.get("GIT_BIN"),
        "M2_REPO_env": os.environ.get("M2_REPO"),
        "M2_REPOS_env": os.environ.get("M2_REPOS"),
        "resolved_mvn": mvn or None,
        "resolved_git": git or None,
        "m2_repos": [str(p) for p in resolve_m2_repos()],
        "gitlab_token_available": default_source != "none",
        "gitlab_token_default_source": default_source,
        "max_batch_projects": MAX_BATCH_PROJECTS,
        "project_pool_size": PROJECT_POOL_SIZE,
        "mvn_version": None,
        "java_version": None,
        "git_version": None,
    }
    if mvn:
        try:
            p = subprocess.run([mvn, "-v"], capture_output=True, text=True, timeout=20)
            info["mvn_version"] = ((p.stdout or "") + (p.stderr or "")).strip()
        except Exception as e:
            info["mvn_version"] = f"ERROR: {e}"
    try:
        p = subprocess.run(["java", "-version"], capture_output=True, text=True, timeout=20)
        info["java_version"] = ((p.stderr or "") + (p.stdout or "")).strip()
    except Exception as e:
        info["java_version"] = f"ERROR: {e}"
    if git:
        try:
            p = subprocess.run([git, "--version"], capture_output=True, text=True, timeout=20)
            info["git_version"] = ((p.stdout or "") + (p.stderr or "")).strip()
        except Exception as e:
            info["git_version"] = f"ERROR: {e}"
    return info


@app.post("/api/v1/analyze/upload")
async def analyze_upload(file: UploadFile = File(...)):
    cleanup_expired()
    if not file.filename or not file.filename.endswith(".zip"):
        raise HTTPException(400, "请上传 .zip 文件")

    content = await file.read()
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, f"文件超过 {MAX_UPLOAD_BYTES // (1024*1024)}MB 限制")

    zip_base = file.filename[:-4] if file.filename.lower().endswith(".zip") else file.filename
    project_name = sanitize_filename(zip_base, fallback="upload")

    task_id = uuid.uuid4().hex
    work_dir = Path(tempfile.mkdtemp(prefix="depgov-"))

    try:
        zip_path = work_dir / "upload.zip"
        zip_path.write_bytes(content)
        extract_dir = work_dir / "src"
        extract_dir.mkdir()
        safe_extract_zip(zip_path, extract_dir)
    except Exception as e:
        shutil.rmtree(work_dir, ignore_errors=True)
        raise HTTPException(400, f"解压失败: {e}")

    with TASKS_LOCK:
        TASKS[task_id] = {
            "task_id": task_id,
            "source": "upload",
            "batch": False,
            "filename": file.filename,
            "status": "pending",
            "created_at": now_iso(),
            "started_at": None,
            "finished_at": None,
            "projects": [{
                "index": 0,
                "repo_url": None,
                "project_name": project_name,
                "status": "pending",
                "started_at": None,
                "finished_at": None,
                "result": None,
                "error": None,
                "warnings": [],
            }],
            "summary": {"total": 1, "success": 0, "failed": 0},
        }

    TASK_EXECUTOR.submit(run_upload_task, task_id, work_dir, extract_dir, project_name)
    return {
        "task_id": task_id,
        "status": "pending",
        "project_count": 1,
        "project_name": project_name,
        "task_url": f"/api/v1/tasks/{task_id}",
        "excel_url": f"/api/v1/tasks/{task_id}/excel",
    }


@app.post("/api/v1/analyze/gitlab")
def analyze_gitlab(
    req: GitLabRequest,
    x_gitlab_token: Optional[str] = Header(default=None, alias="X-GitLab-Token"),
):
    return _create_gitlab_task([req], x_gitlab_token, batch_mode=False)


@app.post("/api/v1/analyze/gitlab/batch")
def analyze_gitlab_batch(
    req: GitLabBatchRequest,
    x_gitlab_token: Optional[str] = Header(default=None, alias="X-GitLab-Token"),
):
    return _create_gitlab_task(req.projects, x_gitlab_token, batch_mode=True)


@app.get("/api/v1/tasks/{task_id}")
def get_task(task_id: str):
    cleanup_expired()
    with TASKS_LOCK:
        task = TASKS.get(task_id)
        if not task:
            raise HTTPException(404, "任务不存在或已过期")

        view = {k: v for k, v in task.items() if k != "token"}

        projects = task.get("projects") or []
        if len(projects) == 1:
            p0 = projects[0]
            view["project_name"] = p0.get("project_name")
            view["error"] = p0.get("error")
            view["warnings"] = p0.get("warnings", [])
            if p0.get("result"):
                legacy = dict(p0["result"])
                legacy["source"] = task.get("source", "gitlab")
                view["result"] = legacy
            else:
                view["result"] = None
            if task.get("filename"):
                view["filename"] = task["filename"]

        if task["status"] in ("success", "partial"):
            view["excel_url"] = f"/api/v1/tasks/{task_id}/excel"
            view["excel_url_direct_only"] = f"/api/v1/tasks/{task_id}/excel?include_transitive=false"
            view["excel_filename"] = f"{build_download_filename(projects)}.xlsx"

        return view


@app.get("/api/v1/tasks/{task_id}/excel")
def export_excel(task_id: str, include_transitive: bool = True):
    cleanup_expired()
    with TASKS_LOCK:
        task = TASKS.get(task_id)
        if not task:
            raise HTTPException(404, "任务不存在或已过期")
        if task["status"] not in ("success", "partial"):
            raise HTTPException(400, f"任务尚未完成，当前状态: {task['status']}")
        projects = list(task.get("projects") or [])

    if not projects:
        raise HTTPException(400, "任务没有可导出的项目")

    content = build_excel(projects, include_transitive=include_transitive)

    base_name = build_download_filename(projects)
    utf8_filename = f"{base_name}.xlsx"
    if len(projects) == 1:
        pname = sanitize_filename(projects[0].get("project_name") or "project", "project")
        ascii_fallback = f"{pname}_ledger.xlsx".replace('"', "_")
    else:
        ascii_fallback = f"ledger_{len(projects)}_projects.xlsx"
    encoded = quote(utf8_filename, safe="")

    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": (
                f'attachment; filename="{ascii_fallback}"; '
                f"filename*=UTF-8''{encoded}"
            ),
            "Content-Length": str(len(content)),
        },
    )
